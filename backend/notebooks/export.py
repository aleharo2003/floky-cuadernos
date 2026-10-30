"""Excel export of the period on screen (HU-027). Dates are dates and amounts are numbers, not text."""

from datetime import datetime
from io import BytesIO

from django.db.models import QuerySet
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from notebooks.models import Entry, Kind, PaymentMethod

MONEY = "#,##0.00;-#,##0.00"
DATE = "DD/MM/YYYY"
BOLD = Font(bold=True)
HEADERS = [
    "Fecha",
    "Tipo",
    "Categoría",
    "Descripción",
    "Medio de pago",
    "Importe",
    "Estado",
    "Espejo",
    "Cargado por",
]


def _autosize(sheet) -> None:
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(width + 2, 10), 60)


def build_workbook(
    *, shop: str, filters: dict, entries: QuerySet[Entry], issued_at: datetime, category: str = ""
) -> bytes:
    book = Workbook()

    detail = book.active
    detail.title = "Movimientos"
    detail.append(HEADERS)
    for cell in detail[1]:
        cell.font = BOLD
    income = expense = 0
    by_category: dict[tuple[str, str], list] = {}
    for e in entries.order_by("date", "created_at"):
        signed = e.amount if e.kind == Kind.INCOME else -e.amount
        detail.append(
            [
                e.date,
                Kind(e.kind).label,
                e.category.name,
                e.description,
                PaymentMethod(e.payment_method).label,
                signed,
                "Anulado" if e.is_voided else "Vigente",
                "Sí" if e.mirror_of_id else "",  # mirrors add up like any other entry (HU-025)
                e.created_by.get_full_name() or e.created_by.get_username(),
            ]
        )
        row = detail.max_row
        detail.cell(row, 1).number_format = DATE
        detail.cell(row, 6).number_format = MONEY
        if e.is_voided:
            continue  # voided entries are listed, but they never add up (HU-027)
        if e.kind == Kind.INCOME:
            income += e.amount
        else:
            expense += e.amount
        group = by_category.setdefault((Kind(e.kind).label, e.category.name), [0, 0])
        group[0] += 1
        group[1] += e.amount
    detail.freeze_panes = "A2"
    _autosize(detail)

    totals = book.create_sheet("Totales")
    period = f"{filters['from']:%d/%m/%Y} al {filters['to']:%d/%m/%Y}"
    for label, value in [
        ("Local", shop),
        ("Período", period),
        ("Emitido", issued_at.strftime("%d/%m/%Y %H:%M")),
        ("Filtros", describe_filters(filters, category) or "Ninguno"),
    ]:
        totals.append([label, value])
    totals.append([])
    totals.append(["Ingresos", income])
    totals.append(["Gastos", expense])
    totals.append(["Saldo", income - expense])
    for row in range(totals.max_row - 2, totals.max_row + 1):
        totals.cell(row, 1).font = BOLD
        totals.cell(row, 2).number_format = MONEY
    totals.append([])
    totals.append(["Tipo", "Categoría", "Movimientos", "Total"])
    for cell in totals[totals.max_row]:
        cell.font = BOLD
    for (kind, category), (count, total) in sorted(by_category.items()):
        totals.append([kind, category, count, total])
        totals.cell(totals.max_row, 4).number_format = MONEY
    _autosize(totals)

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def describe_filters(filters: dict, category: str = "") -> str:
    parts = []
    if filters.get("kind"):
        parts.append(f"tipo {Kind(filters['kind']).label.lower()}")
    if filters.get("category_id"):
        parts.append(f"categoría {category}" if category else "una categoría")
    if filters.get("payment_method"):
        parts.append(f"medio {PaymentMethod(filters['payment_method']).label}")
    if filters.get("q"):
        parts.append(f'texto "{filters["q"]}"')
    if filters.get("voided"):
        parts.append("con anulados (no suman)")
    return ", ".join(parts)
