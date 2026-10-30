"""Filters (HU-028), the year month by month (HU-023) and the Excel export (HU-027)."""

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from accounts.models import Role
from notebooks.export import HEADERS
from notebooks.models import Category, Entry, Kind

pytestmark = pytest.mark.django_db

RANGE = {"from": "2026-10-01", "to": "2026-10-31"}


@pytest.fixture
def book(norte, centro, categories, make_person):
    """A small Local Norte notebook, plus one Centro entry that must never show up."""
    author = make_person("autor")
    sales, suppliers = categories["Local Norte", Kind.INCOME], categories["Local Norte", Kind.EXPENSE]
    rent = Category.objects.create(tenant=norte, kind=Kind.EXPENSE, name="Alquiler")

    def add(day, category, amount, method="cash", description="", voided=False, tenant=norte):
        return Entry.objects.create(
            tenant=tenant,
            category=category,
            kind=category.kind,
            date=day,
            amount=Decimal(amount),
            payment_method=method,
            description=description,
            is_voided=voided,
            void_reason="error" if voided else "",
            created_by=author,
        )

    add(date(2026, 10, 1), sales, "100000.00", "cash", "Jornada viernes")
    add(date(2026, 10, 2), sales, "50000.00", "qr", "Jornada sábado")
    add(date(2026, 10, 2), suppliers, "30000.50", "transfer", "Verdulería")
    add(date(2026, 10, 3), rent, "40000.00", "cash", "Alquiler octubre")
    add(date(2026, 10, 4), suppliers, "9999.00", "cash", "Cargado dos veces", voided=True)
    add(date(2026, 9, 15), sales, "70000.00", "cash", "Septiembre")
    add(date(2026, 10, 2), categories["Local Centro", Kind.INCOME], "1.00", tenant=centro)
    return {"sales": sales, "suppliers": suppliers, "rent": rent}


def descriptions(response):
    assert response.status_code == 200, response.content
    return sorted(e["description"] for e in response.json())


# ---------- filters ----------


@pytest.mark.parametrize(
    ("params", "expected"),
    [
        ({}, ["Alquiler octubre", "Jornada sábado", "Jornada viernes", "Verdulería"]),
        ({"kind": "expense"}, ["Alquiler octubre", "Verdulería"]),
        ({"payment_method": "qr"}, ["Jornada sábado"]),
        ({"q": "VERDU"}, ["Verdulería"]),
        ({"q": "alquiler"}, ["Alquiler octubre"]),  # matches the category name too
        ({"voided": "1", "kind": "expense"}, ["Alquiler octubre", "Cargado dos veces", "Verdulería"]),
    ],
)
def test_list_filters(as_role, book, params, expected):
    api = as_role(Role.VIEWER)

    assert descriptions(api.get("/api/notebooks/entries/", RANGE | params)) == expected


def test_filter_by_category(as_role, book):
    api = as_role(Role.VIEWER)

    response = api.get("/api/notebooks/entries/", RANGE | {"category_id": str(book["rent"].id)})

    assert descriptions(response) == ["Alquiler octubre"]


def test_invalid_filter_is_rejected(as_role, book):
    api = as_role(Role.VIEWER)

    assert api.get("/api/notebooks/entries/", RANGE | {"kind": "otro"}).status_code == 400
    assert api.get("/api/notebooks/entries/", RANGE | {"payment_method": "bitcoin"}).status_code == 400


def test_summary_uses_the_same_filters_and_never_counts_voided(as_role, book):
    api = as_role(Role.VIEWER)

    everything = api.get("/api/notebooks/summary/", RANGE | {"voided": "1"}).json()
    expenses = api.get("/api/notebooks/summary/", RANGE | {"kind": "expense"}).json()

    assert everything["income"] == "150000.00"
    assert everything["expense"] == "70000.50"
    assert everything["balance"] == "79999.50"
    assert expenses["income"] == "0.00"
    assert expenses["expense"] == "70000.50"
    assert {c["category"]: c["total"] for c in expenses["by_category"]} == {
        "Alquiler": "40000.00",
        "Proveedores": "30000.50",
    }


# ---------- year month by month ----------


def test_monthly_report(as_role, book):
    api = as_role(Role.VIEWER)

    response = api.get("/api/notebooks/monthly/", {"year": 2026})

    assert response.status_code == 200
    months = {m["month"]: m for m in response.json()["months"]}
    assert len(months) == 12
    assert months[9] == {"month": 9, "income": "70000.00", "expense": "0.00", "balance": "70000.00"}
    assert months[10] == {"month": 10, "income": "150000.00", "expense": "70000.50", "balance": "79999.50"}
    assert months[1]["balance"] == "0.00"


@pytest.mark.parametrize("year", ["", "abc", "1999"])
def test_monthly_report_needs_a_valid_year(as_role, book, year):
    api = as_role(Role.VIEWER)

    assert api.get("/api/notebooks/monthly/", {"year": year}).status_code == 400


# ---------- Excel ----------


def read_excel(response):
    assert response.status_code == 200, response.content
    return load_workbook(BytesIO(response.content))


def test_excel_has_typed_cells_and_voided_rows_do_not_add_up(as_role, book):
    api = as_role(Role.VIEWER)

    response = api.get("/api/notebooks/export/", RANGE | {"voided": "1"})

    assert response["Content-Disposition"] == 'attachment; filename="cuadernos-20261001-20261031.xlsx"'
    excel = read_excel(response)
    sheet = excel["Movimientos"]
    header = next(sheet.iter_rows(max_row=1, values_only=True))
    assert header == tuple(HEADERS)  # every column has its title
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 5  # the voided one is listed, Centro's is not
    first = rows[0]
    assert first[0].date() == date(2026, 10, 1)  # a real date, not text
    assert first[5] == 100000  # a real number, not text
    voided = next(r for r in rows if r[6] == "Anulado")
    assert voided[3] == "Cargado dos veces"
    assert {r[7] for r in rows} == {None}  # no mirrors yet: the cell stays empty
    expense = next(r for r in rows if r[3] == "Verdulería")
    assert expense[5] == pytest.approx(-30000.50)  # expenses go negative

    totals = {r[0]: r[1] for r in excel["Totales"].iter_rows(values_only=True) if r and r[0]}
    assert totals["Local"] == "Local Norte"
    assert totals["Período"] == "01/10/2026 al 31/10/2026"
    assert totals["Ingresos"] == 150000
    assert totals["Gastos"] == pytest.approx(70000.50)
    assert totals["Saldo"] == pytest.approx(79999.50)


def test_excel_respects_filters(as_role, book):
    api = as_role(Role.VIEWER)

    excel = read_excel(api.get("/api/notebooks/export/", RANGE | {"payment_method": "qr"}))

    rows = list(excel["Movimientos"].iter_rows(min_row=2, values_only=True))
    assert [r[3] for r in rows] == ["Jornada sábado"]
    totals = {r[0]: r[1] for r in excel["Totales"].iter_rows(values_only=True) if r and r[0]}
    assert totals["Filtros"] == "medio QR / Mercado Pago"


def test_excel_leaves_voided_out_unless_asked(as_role, book):
    api = as_role(Role.VIEWER)

    excel = read_excel(api.get("/api/notebooks/export/", RANGE))

    rows = list(excel["Movimientos"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 4
    assert all(r[6] == "Vigente" for r in rows)


def test_export_needs_a_shop_and_a_login(api, book):
    assert api.get("/api/notebooks/export/", RANGE).status_code == 403


def test_excel_names_every_filter_in_words(as_role, book):
    api = as_role(Role.VIEWER)
    params = RANGE | {"kind": "expense", "category_id": str(book["rent"].id), "q": "oct", "voided": "1"}

    excel = read_excel(api.get("/api/notebooks/export/", params))

    totals = {r[0]: r[1] for r in excel["Totales"].iter_rows(values_only=True) if r and r[0]}
    assert totals["Filtros"] == 'tipo gasto, categoría Alquiler, texto "oct", con anulados (no suman)'
